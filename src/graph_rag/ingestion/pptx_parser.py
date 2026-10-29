"""
PPTX Parser module.

Extracts text from Microsoft PowerPoint (.pptx) presentations using python-pptx.
Supports both file paths and in-memory byte streams.
"""

import io
from pathlib import Path
from typing import List, Union
from pptx import Presentation

from graph_rag.ingestion.pdf_parser import PageContent, ParsedDocument


class PptxParser:
    """Extracts text from PPTX slides."""

    def parse(self, source: Union[str, Path, bytes], filename: str = "presentation.pptx") -> ParsedDocument:
        """Parse a .pptx presentation from path or bytes."""
        if isinstance(source, bytes):
            file_stream = io.BytesIO(source)
            prs = Presentation(file_stream)
        else:
            path = Path(source)
            filename = path.name
            prs = Presentation(str(path))

        pages: List[PageContent] = []
        full_text_parts: List[str] = []
        extracted_title = ""

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_lines: List[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_lines.append(text)

            # Slide notes (if present)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    slide_lines.append(f"[Notes: {notes_text}]")

            if slide_lines:
                slide_content = "\n".join(slide_lines)
                pages.append(PageContent(page_number=slide_idx, text=slide_content))
                full_text_parts.append(slide_content)
                if not extracted_title:
                    extracted_title = slide_lines[0]

        if not extracted_title:
            extracted_title = Path(filename).stem.replace("_", " ")

        return ParsedDocument(
            filename=filename,
            title=extracted_title,
            total_pages=len(prs.slides),
            full_text="\n\n".join(full_text_parts),
            pages=pages,
            metadata={"format": "pptx"},
        )


pptx_parser = PptxParser()
