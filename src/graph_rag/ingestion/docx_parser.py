"""
DOCX Parser module.

Extracts text from Microsoft Word (.docx) documents using python-docx.
Supports both file paths and in-memory bytes streams.
"""

import io
from pathlib import Path
from typing import List, Union
from docx import Document

from graph_rag.ingestion.pdf_parser import PageContent, ParsedDocument


class DocxParser:
    """Extracts text and headings from DOCX files."""

    def parse(self, source: Union[str, Path, bytes], filename: str = "document.docx") -> ParsedDocument:
        """Parse a .docx document from path or bytes."""
        if isinstance(source, bytes):
            file_stream = io.BytesIO(source)
            doc = Document(file_stream)
        else:
            path = Path(source)
            filename = path.name
            doc = Document(str(path))

        paragraphs: List[str] = []
        extracted_title = ""

        # Extract paragraphs and detect heading as title
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)
                if not extracted_title and (p.style.name.startswith("Heading") or p.style.name == "Title"):
                    extracted_title = text

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)

        if not extracted_title and paragraphs:
            extracted_title = paragraphs[0][:120]
        if not extracted_title:
            extracted_title = Path(filename).stem.replace("_", " ")

        pages = [PageContent(page_number=1, text=full_text)] if full_text else []

        return ParsedDocument(
            filename=filename,
            title=extracted_title,
            total_pages=1,
            full_text=full_text,
            pages=pages,
            metadata={"format": "docx"},
        )


docx_parser = DocxParser()
